"""
Export service: generates formatted CSV and multi-sheet Excel reports
with financial formatting, metadata, and filter audit details.
"""
from __future__ import annotations

import io
import logging
from datetime import datetime
from typing import Any, Optional

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
import pandas as pd

logger = logging.getLogger(__name__)


class ExportService:
    """Generates professional financial export files."""

    @staticmethod
    def dataframe_to_csv(df: pd.DataFrame) -> bytes:
        """Convert a DataFrame to UTF-8 encoded CSV bytes."""
        return df.to_csv(index=False).encode("utf-8")

    @classmethod
    def generate_excel_report(
        cls,
        sheets: dict[str, pd.DataFrame],
        report_title: str = "Financial Planning Report",
        filters: Optional[dict[str, Any]] = None,
    ) -> bytes:
        """
        Generate a multi-sheet, professionally formatted Excel workbook.

        Parameters:
            sheets: Dictionary mapping sheet name to DataFrame
            report_title: Title shown in metadata
            filters: Dictionary of applied filter settings
        """
        wb = openpyxl.Workbook()
        # Remove default sheet
        wb.remove(wb.active)

        # Header and table styling
        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        regular_font = Font(name="Calibri", size=10)
        bold_font = Font(name="Calibri", size=10, bold=True)
        thin_border = Border(
            left=Side(style="thin", color="E2E8F0"),
            right=Side(style="thin", color="E2E8F0"),
            top=Side(style="thin", color="E2E8F0"),
            bottom=Side(style="thin", color="E2E8F0"),
        )

        for sheet_name, df in sheets.items():
            ws = wb.create_sheet(title=sheet_name[:31])  # Excel sheet title limit is 31 chars

            # Title row
            ws.cell(row=1, column=1, value=report_title).font = Font(name="Calibri", size=14, bold=True, color="0F1729")
            ws.cell(row=2, column=1, value=f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}").font = Font(name="Calibri", size=9, italic=True, color="64748B")

            # Column headers on row 4
            start_row = 4
            for col_idx, col_name in enumerate(df.columns, start=1):
                cell = ws.cell(row=start_row, column=col_idx, value=col_name)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = thin_border

            # Data rows
            for row_idx, row_data in enumerate(df.itertuples(index=False), start=start_row + 1):
                for col_idx, val in enumerate(row_data, start=1):
                    cell = ws.cell(row=row_idx, column=col_idx, value=val)
                    cell.font = regular_font
                    cell.border = thin_border

                    # Financial numeric formatting
                    if isinstance(val, (int, float)):
                        col_name = str(df.columns[col_idx - 1]).lower()
                        if "%" in col_name or "margin" in col_name:
                            cell.number_format = "0.00%"
                            cell.alignment = Alignment(horizontal="right")
                        elif "€" in col_name or "amount" in col_name or "revenue" in col_name or "profit" in col_name or "variance" in col_name:
                            cell.number_format = "#,##0.00"
                            cell.alignment = Alignment(horizontal="right")
                        else:
                            cell.alignment = Alignment(horizontal="right")
                    else:
                        cell.alignment = Alignment(horizontal="left")

            # Auto-adjust column widths
            for col in ws.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    if cell.row < start_row:
                        continue
                    if cell.value:
                        max_len = max(max_len, len(str(cell.value)))
                ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        # Metadata Sheet
        meta_ws = wb.create_sheet(title="Report Info")
        meta_ws.cell(row=1, column=1, value="Report Metadata").font = Font(name="Calibri", size=14, bold=True)
        meta_ws.cell(row=3, column=1, value="Company:").font = bold_font
        meta_ws.cell(row=3, column=2, value="Northstar Digital Services GmbH")
        meta_ws.cell(row=4, column=1, value="Report Title:").font = bold_font
        meta_ws.cell(row=4, column=2, value=report_title)
        meta_ws.cell(row=5, column=1, value="Generated At:").font = bold_font
        meta_ws.cell(row=5, column=2, value=datetime.now().strftime("%Y-%m-%d %H:%M:%S"))

        if filters:
            meta_ws.cell(row=7, column=1, value="Applied Filters").font = Font(name="Calibri", size=12, bold=True)
            r = 8
            for k, v in filters.items():
                meta_ws.cell(row=r, column=1, value=f"{k}:").font = bold_font
                meta_ws.cell(row=r, column=2, value=str(v))
                r += 1

        output = io.BytesIO()
        wb.save(output)
        return output.getvalue()
